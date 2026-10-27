"""Text extraction from PDF and DOCX.

Text-based files only. A scanned resume produces no text, and rather than guess we
raise :class:`NoTextFound` so the UI can offer the manual fallback with the exact
copy from the plan. OCR is a v1.1 feature.
"""

from __future__ import annotations

import asyncio
import io
import re
from enum import StrEnum

from app.core.config import settings
from app.core.errors import PayloadTooLarge, UnsupportedMedia

# Two pages of dense CV is ~8k chars. 60k is generous and bounds every LLM prompt.
MAX_CHARS = 60_000

PDF_MAGIC = b"%PDF-"
ZIP_MAGIC = b"PK\x03\x04"

CONTROL_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")


class DocumentKind(StrEnum):
    PDF = "pdf"
    DOCX = "docx"


class NoTextFound(Exception):
    """The file is a valid document but carries no extractable text."""


def detect_kind(filename: str | None, content_type: str | None, data: bytes) -> DocumentKind:
    """Trust the bytes, not the header: uploads are attacker-controlled."""
    if len(data) == 0:
        raise UnsupportedMedia("That file is empty.")

    if len(data) > settings.upload_max_mb * 1024 * 1024:
        raise PayloadTooLarge(f"Resumes must be under {settings.upload_max_mb} MB.")

    if data.startswith(PDF_MAGIC):
        return DocumentKind.PDF
    if data.startswith(ZIP_MAGIC):
        # A .docx is a zip; a plain .zip must not slip through.
        if b"word/document.xml" in data:
            return DocumentKind.DOCX
        raise UnsupportedMedia("That archive is not a Word document.")

    raise UnsupportedMedia("Upload a text-based PDF or DOCX file.")


def _clean(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = CONTROL_CHARS.sub(" ", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as exc:
        raise NoTextFound("We couldn't open that PDF.") from exc

    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as exc:
            raise NoTextFound("That PDF is password protected.") from exc

    chunks: list[str] = []
    for page in reader.pages:
        try:
            chunks.append(page.extract_text() or "")
        except Exception:
            # One unreadable page should not lose the rest of the resume.
            continue

    return _clean("\n".join(chunks))


def _extract_docx(data: bytes) -> str:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise NoTextFound("We couldn't open that DOCX file.") from exc

    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    return _clean("\n".join(parts))


def extract_text_sync(data: bytes, kind: DocumentKind) -> str:
    text = _extract_pdf(data) if kind is DocumentKind.PDF else _extract_docx(data)
    if len(text) < 40:
        raise NoTextFound("No readable text in that file.")
    return text[:MAX_CHARS]


async def extract_text(data: bytes, kind: DocumentKind) -> str:
    """Parsing is CPU-bound, so keep it off the event loop."""
    return await asyncio.to_thread(extract_text_sync, data, kind)


NO_TEXT_MESSAGE = (
    "We couldn't read this file. Upload a text-based PDF or DOCX, or fill your details in manually."
)
